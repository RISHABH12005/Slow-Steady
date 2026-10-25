from io import BytesIO
from fastapi.testclient import TestClient
from docx import Document
from pptx import Presentation
import app as module

client = TestClient(module.app)

def sample_data():
    return {
        "document_type": "internship",
        "project_title": "Example Project",
        "student_name": "Test Student",
        "enrollment": "TEST123",
        "organization": "Example Org",
        "guide": "Example Guide",
        "department": "Computer Science and Engineering",
        "year": "2026",
        "overview": "A test overview.",
        "work": "A test implementation description.",
        "results": "Test results are illustrative.",
        "contribution": "Test contribution.",
        "conclusion": "Test conclusion.",
        "references": "[1] Example reference; replace with verified citation."
    }

def test_home_is_dark_mode_and_lists_profiles():
    r = client.get("/")
    assert r.status_code == 200
    assert "color-scheme:dark" in r.text
    assert "Minor Project I Report" in r.text
    assert "Major Project Presentation" in r.text

def test_health_and_profile_api():
    assert client.get("/health").json()["status"] == "ok"
    profiles = client.get("/api/profiles").json()
    assert {"internship", "minor1", "minor2", "major1", "major2"} <= set(profiles["reports"])
    assert sum(x["slides"] for x in profiles["major_ppt"]) == 13

def test_report_docx_is_valid_and_has_guideline_formatting():
    data = sample_data()
    doc = module.build_report(data)
    parsed = Document(BytesIO(doc.getvalue()))
    assert any("Example Project" in p.text for p in parsed.paragraphs)
    assert parsed.sections[0].left_margin.inches == 1.5
    assert parsed.sections[0].right_margin.inches == 1.0
    assert parsed.styles["Normal"].font.name == "Times New Roman"
    assert parsed.styles["Normal"].font.size.pt == 12
    assert any("REFERENCES" in p.text for p in parsed.paragraphs)

def test_all_report_profiles_generate_docx():
    for profile in ("internship", "minor1", "minor2", "major1", "major2"):
        data = sample_data()
        data["document_type"] = profile
        doc = Document(BytesIO(module.build_report(data).getvalue()))
        assert len(doc.paragraphs) > 10
        assert any(p.style.name.startswith("Heading") for p in doc.paragraphs)

def test_pptx_is_valid_and_has_required_nine_sections():
    data = sample_data()
    prs = Presentation(BytesIO(module.build_ppt(data).getvalue()))
    assert len(prs.slides) == 13
    text = "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
    for heading in ("Problem Statement", "Literature Review", "Work Plan", "Individual Contribution", "Work Done So Far", "Work to be Done", "Conclusion", "References"):
        assert heading in text
    assert "AI-generated images" not in text

def test_generate_route_returns_docx_and_pptx():
    data = sample_data()
    r = client.post("/generate", data=data)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    data["document_type"] = "major_ppt"
    r = client.post("/generate", data=data)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.presentationml.presentation")
    assert Presentation(BytesIO(r.content)).slides

def test_validator_warns_when_content_missing():
    r = client.post("/api/validate", data={"document_type":"major_ppt","project_title":"X","student_name":"Y"})
    assert r.status_code == 200
    assert any(i["severity"]=="warning" for i in r.json()["issues"])
