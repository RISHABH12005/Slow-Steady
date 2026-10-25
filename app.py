from __future__ import annotations
from io import BytesIO
from datetime import date
from html import escape
from fastapi import FastAPI, Form
from fastapi.responses import HTMLResponse, StreamingResponse, JSONResponse
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.dml.color import RGBColor

app = FastAPI(title="Slow-Steady", description="JUET academic report and presentation generator", version="0.2.0")

REPORT_PROFILES = {
    "internship": {"label":"Summer Internship Report", "chapters":[
        ("1. INTRODUCTION", "Describe the specific project, problem, objectives and scope."),
        ("2. DESCRIPTION OF ORGANIZATION", "Use verified facts about the organization and team."),
        ("3. DESCRIPTION OF WORK", "Describe internship activities in logical sections."),
        ("4. SYSTEM ANALYSIS AND DESIGN", "Application-based: architecture, methodology and DFD/UML. Research-based: Proposed Methodology."),
        ("5. IMPLEMENTATION AND TESTING", "Explain implementation and tests with actual results; do not include source code."),
        ("6. RESULTS AND DISCUSSION", "Present verified results and interpretation. Research-based work should compare existing work."),
        ("7. CONCLUSION", "Summarize work, outcomes, limitations and learning."),
        ("8. REFERENCES", "List sources actually cited; verify details and citation order."),
        ("9. APPENDICES", "Supplementary material labelled Appendix A, Appendix B, etc.")
    ]},
    "minor1": {"label":"Minor Project I Report", "chapters":[
        ("1. INTRODUCTION","Problem context, motivation, objectives and scope."),
        ("2. PROBLEM STATEMENT","Define the problem and intended users."),
        ("3. LITERATURE REVIEW","Summarize relevant prior work with verified citations."),
        ("4. PROPOSED METHODOLOGY","Describe the planned approach and rationale."),
        ("5. WORK PLAN","Milestones, timeline and deliverables."),
        ("6. EXPECTED OUTCOMES","State intended outcomes, not fabricated results."),
        ("7. CONCLUSION","Summarize the proposal and next steps."),
        ("8. REFERENCES","List sources actually cited.")
    ]},
    "minor2": {"label":"Minor Project II Report", "chapters":[
        ("1. INTRODUCTION","Project context, objectives and scope."),
        ("2. LITERATURE REVIEW","Prior work and identified gap."),
        ("3. SYSTEM ANALYSIS AND DESIGN","Requirements, architecture, design and diagrams."),
        ("4. IMPLEMENTATION / METHODOLOGY","Describe actual work completed; adapt for research projects."),
        ("5. TESTING / ANALYSIS","Test cases, experiments and verified outcomes."),
        ("6. RESULTS AND DISCUSSION","Interpret results and note limitations."),
        ("7. CONCLUSION AND FUTURE WORK","Summarize progress and next steps."),
        ("8. REFERENCES","List sources actually cited.")
    ]},
    "major1": {"label":"Major Project I Report", "chapters":[
        ("1. INTRODUCTION","Context, motivation, problem and objectives."),
        ("2. PROBLEM STATEMENT","Specific problem, scope and constraints."),
        ("3. LITERATURE REVIEW","Prior work and research/engineering gap."),
        ("4. REQUIREMENTS AND METHODOLOGY","Requirements, proposed method, design and plan."),
        ("5. WORK COMPLETED","Evidence-backed progress and contribution of each member."),
        ("6. WORK PLAN","Remaining milestones and timeline."),
        ("7. CONCLUSION","Current status, risks and next steps."),
        ("8. REFERENCES","Verified sources in the department-required style.")
    ]},
    "major2": {"label":"Major Project II Report", "chapters":[
        ("1. INTRODUCTION","Context, problem, objectives and scope."),
        ("2. LITERATURE REVIEW","Prior work and gap."),
        ("3. SYSTEM ANALYSIS AND DESIGN","Architecture, methodology and relevant diagrams."),
        ("4. IMPLEMENTATION / EXPERIMENTS","Actual implementation or research experiments."),
        ("5. TESTING AND EVALUATION","Methods, test cases, measurements and results."),
        ("6. RESULTS AND DISCUSSION","Interpretation, comparisons and limitations."),
        ("7. CONCLUSION AND FUTURE WORK","Conclusions supported by evidence."),
        ("8. REFERENCES","Verified sources in the required style."),
        ("9. APPENDICES","Supplementary material.")
    ]}
}
PPT_SECTIONS = [
 ("Introduction",1,"Project context, motivation and objective."),
 ("Problem Statement",1,"Specific problem, users, scope and constraints."),
 ("Literature Review",2,"Prior work and the gap; use 1–2 slides."),
 ("Work Plan",1,"Timeline and milestones aligned with weekly reports."),
 ("Individual Contribution",1,"Clearly identify each member's actual contribution."),
 ("Work Done So Far",4,"Use 3–4 slides with evidence. Application-based: design, implementation, testing. Research-based: methodology, experiments, analysis."),
 ("Work to be Done",1,"Remaining tasks, milestones and risks."),
 ("Conclusion",1,"Current status and evidence-backed conclusion."),
 ("References",1,"References in IEEE format.")
]

CSS = """
:root{color-scheme:dark;--bg:#090d16;--panel:#111827;--panel2:#151f30;--line:#273449;--text:#eef3fb;--muted:#9baac0;--blue:#83aaff;--green:#75ddb0;font-family:Inter,ui-sans-serif,system-ui,Segoe UI,Arial,sans-serif}
*{box-sizing:border-box}body{margin:0;background:radial-gradient(ellipse at 80% -10%,#172949 0,transparent 42%),var(--bg);color:var(--text)}
.wrap{max-width:1120px;margin:auto;padding:28px 20px 60px}.top{display:flex;justify-content:space-between;align-items:center;gap:16px;margin-bottom:28px}.brand{font-weight:800;letter-spacing:-.04em;font-size:20px}.brand span{color:var(--blue)}.tag{border:1px solid var(--line);border-radius:999px;padding:7px 11px;color:var(--muted);font-size:12px}
.hero{padding:34px;border:1px solid #2a3b56;border-radius:22px;background:linear-gradient(120deg,#121e31,#111827 65%);box-shadow:0 22px 60px #0003}.eyebrow{font-size:12px;text-transform:uppercase;letter-spacing:.16em;color:var(--blue);font-weight:750}h1{font-size:clamp(34px,5vw,56px);line-height:1.04;letter-spacing:-.05em;margin:15px 0}h2{font-size:21px;letter-spacing:-.02em;margin:0 0 10px}h3{font-size:15px;margin:0 0 8px}.hero p,.muted{color:var(--muted);line-height:1.65;max-width:760px}.stats{display:grid;grid-template-columns:repeat(3,1fr);gap:12px;margin-top:24px}.stat{padding:14px;background:#0b1220;border:1px solid var(--line);border-radius:12px}.stat strong{display:block;font-size:18px}.stat span{font-size:12px;color:var(--muted)}
.tabs{display:flex;gap:8px;flex-wrap:wrap;margin:22px 0}.tab{border:1px solid var(--line);border-radius:9px;padding:10px 13px;color:var(--muted);text-decoration:none}.tab.active,.tab:hover{background:#1a2c48;color:var(--text);border-color:#46638c}
.layout{display:grid;grid-template-columns:minmax(0,1.5fr) minmax(260px,.8fr);gap:18px;align-items:start}.panel{background:color-mix(in srgb,var(--panel) 95%,transparent);border:1px solid var(--line);border-radius:16px;padding:22px;margin-bottom:18px;min-width:0}.grid{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:15px}.full{grid-column:1/-1}label{display:block;font-weight:650;font-size:12px;color:#d9e3f1;margin-bottom:7px}input,textarea,select{width:100%;padding:12px 13px;border:1px solid #33435b;border-radius:9px;background:#0b1220;color:var(--text);font:inherit;outline:none}input:focus,textarea:focus,select:focus{border-color:var(--blue);box-shadow:0 0 0 3px #83aaff18}textarea{min-height:88px;resize:vertical}button{font:inherit;cursor:pointer}.primary{background:var(--blue);color:#07101f;border:0;padding:13px 17px;border-radius:10px;font-weight:800}.secondary{background:#182337;color:var(--text);border:1px solid var(--line);padding:12px 15px;border-radius:10px;font-weight:700}.primary:hover{filter:brightness(1.08)}.row{display:flex;gap:10px;flex-wrap:wrap;align-items:center}.muted{font-size:13px}.section{padding-top:18px;margin-top:18px;border-top:1px solid var(--line)}.list{display:grid;gap:10px}.item{padding:12px;background:#0b1220;border:1px solid var(--line);border-radius:10px}.item p{font-size:13px;color:var(--muted);margin:5px 0 0;line-height:1.5}.check{color:var(--green)}footer{color:#7889a1;font-size:12px;padding:12px 0}
.preview-shell{position:sticky;top:12px}.preview-page{background:#fff;color:#111;border-radius:3px;box-shadow:0 12px 34px #0006;padding:34px 28px;min-height:410px;aspect-ratio:210/297;font-family:"Times New Roman",serif;font-size:10px;overflow:hidden;border:1px solid #d4dbe5}.preview-page .cover-title{text-align:center;text-transform:uppercase;font-weight:bold;font-size:14px;line-height:1.25;margin:18px 0 22px}.preview-page .cover-label{text-align:center;margin:12px 0}.preview-page .cover-name{text-align:center;font-size:12px;font-weight:bold;margin:14px 0}.preview-page .cover-university{text-align:center;font-weight:bold;margin-top:28px}.preview-page .preview-rule{border:0;border-top:1px solid #ddd;margin:20px 0}.preview-page .preview-chapter{font-size:12px;font-weight:bold;margin-top:12px}.preview-page .preview-copy{font-size:10px;line-height:1.35;white-space:pre-wrap;overflow-wrap:anywhere}.preview-status{font-size:12px;color:var(--muted);margin:9px 0 0}.preview-page .placeholder{color:#667085;font-style:italic}.preview-page .preview-page-number{text-align:center;margin-top:28px;color:#555}.@media(max-width:780px){.layout{grid-template-columns:1fr}.hero{padding:24px}.stats{grid-template-columns:1fr 1fr}}@media(max-width:520px){.wrap{padding:16px 12px 40px}.grid{grid-template-columns:1fr}.full{grid-column:auto}.stats{grid-template-columns:1fr}.top{align-items:flex-start}}
/* Professional UI refinement */
:root{--bg:#080b12;--panel:#101722;--panel2:#141e2b;--line:#263244;--text:#f3f6fb;--muted:#9ba8ba;--blue:#9bbcff;--green:#7ddbb5}
body{background:linear-gradient(180deg,#0b1220 0,#080b12 420px);font-size:14px;letter-spacing:-.01em}
.wrap{max-width:1240px;padding:26px 28px 64px}
.top{padding:8px 2px 20px;border-bottom:1px solid #1f2a3a;margin-bottom:26px}
.brand{font-size:22px;letter-spacing:-.055em}.brand span{color:#8eb4ff}
.tag{background:#0e1623;border-color:#2b394d;font-size:11px;letter-spacing:.04em}
.hero{position:relative;overflow:hidden;padding:clamp(30px,5vw,58px);border-radius:24px;border-color:#2a3a50;background:linear-gradient(115deg,#142239 0%,#101827 55%,#10151f 100%);box-shadow:0 24px 70px #0004}
.hero:after{content:"";position:absolute;right:-90px;top:-140px;width:390px;height:390px;border:1px solid #8eb4ff22;border-radius:50%;box-shadow:0 0 0 45px #8eb4ff08,0 0 0 90px #8eb4ff05;pointer-events:none}
.eyebrow{letter-spacing:.18em;font-size:11px}
h1{font-size:clamp(38px,5vw,64px);line-height:1.02;max-width:800px;letter-spacing:-.065em;margin:18px 0}
.hero p{font-size:15px;max-width:650px}
.stats{max-width:700px;gap:10px;margin-top:30px}
.stat{background:#0b1320b8;border-color:#34445b;padding:16px;border-radius:13px}
.stat strong{font-size:17px;letter-spacing:-.03em}.stat span{display:block;margin-top:5px}
.tabs{gap:7px;margin:20px 0 18px}.tab{font-size:12px;padding:10px 14px;background:#0d1420;border-color:#263449;transition:background .18s,border-color .18s}
.tab.active{background:#1a2b43;border-color:#49648b;color:#fff}
.layout{grid-template-columns:minmax(0,1.12fr) minmax(320px,.88fr);gap:20px}
.panel{padding:25px;border-radius:17px;background:linear-gradient(180deg,#111a27,#0e1520);border-color:#263449;box-shadow:0 10px 30px #00000012}
.panel>h2{font-size:20px;letter-spacing:-.035em}
label{font-size:12px;letter-spacing:.01em;color:#dce5f2;margin-bottom:8px}
input,textarea,select{background:#090f19;border-color:#2c3b50;border-radius:10px;padding:13px 14px;font-size:13px;transition:border-color .16s,box-shadow .16s}
input::placeholder,textarea::placeholder{color:#6f7d91}
input:focus,textarea:focus,select:focus{border-color:#8eb4ff;box-shadow:0 0 0 3px #8eb4ff15}
.grid{gap:17px}.primary,.secondary{border-radius:10px;transition:transform .15s,filter .15s,background .15s}
.primary{background:#a6c2ff;color:#08111f;padding:13px 18px;box-shadow:0 5px 18px #8eb4ff1a}
.primary:hover,.secondary:hover{transform:translateY(-1px)}
.secondary{background:#141f2e;border-color:#34455c}.muted{line-height:1.65}
.item{border-radius:12px;border-color:#26364b;background:#0b121d}
.preview-shell{top:18px}.preview-page{border-radius:4px;box-shadow:0 20px 55px #0007;min-height:500px;padding:clamp(28px,4vw,44px) clamp(24px,3vw,36px)}
.preview-status{line-height:1.5}
.section{border-color:#263449}
footer{padding-top:20px;border-top:1px solid #1c2736;margin-top:20px}
@media(max-width:900px){.layout{grid-template-columns:1fr}.preview-shell{position:static;max-width:650px;margin:0 auto;width:100%}.preview-page{min-height:560px}}
@media(max-width:520px){.wrap{padding:14px 12px 36px}.top{flex-direction:column;gap:10px}.hero{padding:25px 20px}.hero:after{opacity:.35}.stats{grid-template-columns:1fr}.panel{padding:18px}.preview-page{min-height:480px}}

"""
PAGE = """<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Slow-Steady · Document Studio</title><style>""" + CSS + """
:root{--accent:#a5c2ff;--accent2:#d1e0ff;--surface:#111a27;--surface2:#0c131e}
body{background:#080c13;font-size:14px}.wrap{max-width:1380px;padding:24px 30px 60px}
.top{margin-bottom:28px;padding-bottom:18px}.brand{font-size:21px}.brand span{color:var(--accent)}
.kicker{font-size:11px;text-transform:uppercase;letter-spacing:.16em;color:#a6bbdc;font-weight:800}
.dashboard-intro{max-width:760px;margin:18px 0 28px}.dashboard-intro h1{font-size:clamp(36px,5vw,54px);letter-spacing:-.055em;line-height:1.05;margin:12px 0}.dashboard-intro p{color:var(--muted);font-size:15px;line-height:1.7}
.screen{display:none}.screen.active{display:block}.choice-grid{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:14px}
.choice{display:flex;flex-direction:column;gap:14px;text-align:left;color:var(--text);padding:21px;min-height:190px;background:linear-gradient(145deg,#141f2e,#0e1520);border:1px solid #2a394d;border-radius:16px;transition:transform .18s,border-color .18s,background .18s;cursor:pointer}
.choice:hover{transform:translateY(-3px);border-color:#6b8dbd;background:linear-gradient(145deg,#1b2b40,#101a29)}
.choice:focus-visible{outline:2px solid var(--accent);outline-offset:3px}.choice-icon{display:grid;place-items:center;width:40px;height:40px;background:#243751;border:1px solid #3b5271;border-radius:12px;color:var(--accent);font-weight:800;font-size:16px}
.choice h2{font-size:16px;margin:0;letter-spacing:-.025em}.choice p{font-size:12px;line-height:1.55;color:var(--muted);margin:7px 0 0}.choice-bottom{display:flex;justify-content:space-between;align-items:center;margin-top:auto;color:#a8c3f3;font-size:12px;font-weight:700}.choice-type{font-size:10px;text-transform:uppercase;letter-spacing:.12em;color:#8e9fb7}
.workspace-head{display:flex;align-items:center;justify-content:space-between;gap:16px;margin-bottom:18px}.workspace-head h1{font-size:26px;letter-spacing:-.04em;margin:5px 0}.back-btn{background:#111b29;color:#d6e4f9;border:1px solid #2d4059;border-radius:10px;padding:11px 14px;font-weight:700}.back-btn:hover{background:#19283b}
.workspace-layout{display:grid;grid-template-columns:minmax(0,1fr) minmax(360px,.9fr);gap:18px;align-items:start}
.workspace-form{max-height:none}.form-scroll{display:block}.workspace-preview{position:sticky;top:12px}.preview-toolbar{display:flex;align-items:center;justify-content:space-between;gap:10px;margin-bottom:14px}.preview-controls{display:flex;align-items:center;justify-content:space-between;gap:8px;margin:12px 0 8px}.preview-controls .secondary{padding:8px 11px;font-size:12px}.preview-controls span{font-size:12px;color:#c4d0e0;font-variant-numeric:tabular-nums}.preview-hint{font-size:12px;color:var(--muted);line-height:1.5;margin:0 0 12px}.preview-pages .preview-page{display:none}.preview-pages .preview-page.active{display:block}.preview-page{max-width:520px;min-height:0;height:auto;aspect-ratio:210/297;margin:0 auto;padding:10% 9%}.workspace-preview .preview-pages.presentation .preview-page{aspect-ratio:16/9;max-width:640px;padding:8%}.live-dot{display:inline-flex;align-items:center;gap:7px;font-size:11px;color:#8ee0bc}.live-dot:before{content:"";width:7px;height:7px;background:#69d7a8;border-radius:50%;box-shadow:0 0 0 3px #69d7a81c}
.preview-pages{display:grid;gap:12px;max-height:none;overflow:visible;padding:0}.preview-page{width:100%;max-width:520px;min-height:0;aspect-ratio:210/297;height:auto;margin:0 auto;padding:10% 9%;background:#fff;color:#111;border:1px solid #d4dbe5;border-radius:2px;box-shadow:0 14px 35px #0007;overflow:hidden;position:relative;scroll-snap-align:start;font-family:"Times New Roman",serif;font-size:clamp(8px,.8vw,11px)}.preview-page .cover-title{text-align:center;text-transform:uppercase;font-weight:bold;font-size:1.45em;line-height:1.25;margin:10% 0 12%}.preview-page .cover-label{text-align:center;margin:5% 0}.preview-page .cover-name{text-align:center;font-size:1.15em;font-weight:bold;margin:6% 0}.preview-page .cover-university{text-align:center;font-weight:bold;margin-top:16%}.preview-page .preview-rule{border:0;border-top:1px solid #ddd;margin:8% 0}.preview-page .preview-chapter{font-size:1.2em;font-weight:bold;margin:7% 0 3%}.preview-page .preview-copy{font-size:1em;line-height:1.5;white-space:pre-wrap;overflow-wrap:anywhere}.preview-page .preview-page-number{position:absolute;bottom:5%;left:0;right:0;text-align:center;color:#555}.preview-page .page-running-title{font-size:.85em;color:#666;text-align:center;border-bottom:1px solid #ddd;padding-bottom:4%;margin-bottom:6%}.preview-page .page-content{max-height:75%;overflow:hidden;white-space:pre-wrap;overflow-wrap:anywhere;line-height:1.5}.preview-page .page-number-label{position:absolute;top:3%;right:6%;font:10px ui-sans-serif,system-ui;color:#8a94a3}.workspace-preview .preview-pages.presentation .preview-page{aspect-ratio:16/9;max-width:640px;padding:8%;font-family:Arial,sans-serif}.workspace-preview .preview-pages.presentation .cover-title{text-transform:none;font-size:1.8em;margin:12% 0}.preview-placeholder{font-family:inherit;color:var(--muted);font-size:13px;line-height:1.6;padding:14px;background:#0c1420;border:1px solid #27364a;border-radius:10px;margin-bottom:14px}
.format-note{font-size:12px;color:#99a8bd;line-height:1.6;margin-top:12px}.field-help{font-size:11px;color:#8596ac;margin-top:6px;line-height:1.45}
.action-bar{display:flex;gap:10px;flex-wrap:wrap;margin-top:22px;padding-top:18px;border-top:1px solid #273449}.hidden{display:none!important}
@media(max-width:1080px){.choice-grid{grid-template-columns:repeat(2,minmax(0,1fr))}.workspace-layout{grid-template-columns:minmax(0,1fr) minmax(300px,.85fr)}}
@media(max-width:760px){.wrap{padding:16px 13px 40px}.choice-grid{grid-template-columns:1fr}.choice{min-height:0;padding:18px}.workspace-layout{grid-template-columns:1fr}.workspace-preview{position:static}.workspace-preview .preview-page{min-height:540px}.workspace-head{align-items:flex-start;flex-direction:column}.workspace-head h1{font-size:23px}.dashboard-intro h1{font-size:37px}.panel{padding:18px}.workspace-form{order:1}.workspace-preview{order:2}}
</style></head><body><main class="wrap">
<header class="top"><div class="brand">slow<span>·</span>steady <span style="color:#7d8ca2;font-size:12px;font-weight:500;letter-spacing:0">/ DOCUMENT STUDIO</span></div><div class="tag">JUET ACADEMIC WORKSPACE</div></header>
<section id="dashboard" class="screen active" aria-labelledby="dashTitle">
<div class="dashboard-intro"><div class="kicker">Workspace / 01</div><h1 id="dashTitle">What are you working on?</h1><p>Choose a document to get started. Your editor and live page preview will open together in the next step, so you can work and review the layout side by side.</p></div>
<div class="choice-grid">
<button type="button" class="choice" data-type="internship"><span class="choice-icon">01</span><div><h2>Summer Internship Report</h2><p>Organization profile, work completed, implementation, results and references.</p></div><div class="choice-bottom"><span class="choice-type">Report · DOCX</span><span>Continue →</span></div></button>
<button type="button" class="choice" data-type="minor1"><span class="choice-icon">02</span><div><h2>Minor Project I Report</h2><p>Problem definition, literature review, methodology and project plan.</p></div><div class="choice-bottom"><span class="choice-type">Report · DOCX</span><span>Continue →</span></div></button>
<button type="button" class="choice" data-type="minor2"><span class="choice-icon">03</span><div><h2>Minor Project II Report</h2><p>Project implementation or research progress, analysis and results.</p></div><div class="choice-bottom"><span class="choice-type">Report · DOCX</span><span>Continue →</span></div></button>
<button type="button" class="choice" data-type="major1"><span class="choice-icon">04</span><div><h2>Major Project I Report</h2><p>Requirements, methodology, work completed and the next milestones.</p></div><div class="choice-bottom"><span class="choice-type">Report · DOCX</span><span>Continue →</span></div></button>
<button type="button" class="choice" data-type="major2"><span class="choice-icon">05</span><div><h2>Major Project II Report</h2><p>Final design, implementation, testing, evaluation and conclusions.</p></div><div class="choice-bottom"><span class="choice-type">Report · DOCX</span><span>Continue →</span></div></button>
<div class="sr-only-profile-labels" aria-hidden="true" style="position:absolute;width:1px;height:1px;overflow:hidden;clip-path:inset(50%)"><span>Major Project Presentation</span></div>
<button type="button" class="choice" data-type="major_ppt"><span class="choice-icon">06</span><div><h2>Major Project Presentation</h2><p>Required major-project review sections, individual contribution, progress evidence and IEEE references.</p></div><div class="choice-bottom"><span class="choice-type">Presentation · PPTX</span><span>Continue →</span></div></button>
<button type="button" class="choice" data-type="minor_ppt"><span class="choice-icon">07</span><div><h2>Major & Minor Project PPT</h2><p>Project problem, approach, progress, key evidence and next steps.</p></div><div class="choice-bottom"><span class="choice-type">Presentation · PPTX</span><span>Continue →</span></div></button>
<button type="button" class="choice" data-type="internship_ppt"><span class="choice-icon">08</span><div><h2>Summer Internship Report PPT</h2><p>Internship overview, organization, work completed, outcomes and learning.</p></div><div class="choice-bottom"><span class="choice-type">Presentation · PPTX</span><span>Continue →</span></div></button>
</div>
<p class="muted" style="margin-top:22px">All exported files are editable drafts. Verify your department's current format and review the final file before submission.</p>
</section>
<section id="workspace" class="screen" aria-labelledby="workspaceTitle">
<div class="workspace-head"><div><div class="kicker">Workspace / 02</div><h1 id="workspaceTitle">Build your document</h1><div id="workspaceSubtitle" class="muted">Enter project details on the left. The preview updates as you type.</div></div><button type="button" id="backToDashboard" class="back-btn">← All document types</button></div>
<div class="workspace-layout">
<form id="builderForm" class="panel workspace-form" method="post" action="/generate">
<h2>Document details</h2><p class="muted">Add verified project information. Required fields are marked with an asterisk.</p>
<input type="hidden" name="document_type" id="document_type" value="internship">
<div class="grid">
<div class="full"><label>Project title *</label><input name="project_title" required placeholder="A specific, meaningful project title"></div>
<div><label>Student / presenter *</label><input name="student_name" required placeholder="Full name"></div><div><label>Enrollment number</label><input name="enrollment" placeholder="Enrollment number"></div>
<div id="organizationField"><label>Organization / company</label><input name="organization" placeholder="Internship organization"></div><div><label>Supervisor / project guide</label><input name="guide" placeholder="Guide or supervisor"></div>
<div><label>Department</label><input name="department" value="Computer Science and Engineering"></div><div><label>Academic year</label><input name="year" type="number" value="2026"></div>
<div class="full"><label>Project overview / problem statement</label><textarea name="overview" placeholder="Problem, intended users, objectives and scope"></textarea></div>
<div class="full"><label>Work completed / methodology</label><textarea name="work" placeholder="Actual work completed, design, implementation or research methodology"></textarea></div>
<div class="full"><label>Results / analysis</label><textarea name="results" placeholder="Verified findings, test results, measurements and limitations"></textarea></div>
<div class="full"><label>Individual contribution / work plan</label><textarea name="contribution" placeholder="Actual contribution by each member and/or milestones"></textarea></div>
<div class="full"><label>Conclusion / future work</label><textarea name="conclusion" placeholder="Evidence-backed conclusions and remaining work"></textarea></div>
<div class="full"><label>References (one per line)</label><textarea name="references" placeholder="Paste verified references; include only sources you used"></textarea></div>
</div>
<div class="action-bar"><button class="primary" type="submit">Generate editable file ↗</button><button class="secondary" type="reset">Clear fields</button></div>
<p class="muted">The right-hand preview is a live layout guide. Always open the exported DOCX/PPTX to verify exact pagination and content.</p>
</form>
<aside class="panel workspace-preview"><div class="preview-toolbar"><div><h2 style="margin:0">Live preview</h2><div class="muted" style="margin-top:5px">Your document, page by page</div></div><span class="live-dot">LIVE</span></div><div class="preview-controls"><button type="button" class="secondary" id="prevPreviewPage">← Previous</button><span id="previewPageIndicator">Page 1 of 5</span><button type="button" class="secondary" id="nextPreviewPage">Next →</button></div><p class="preview-hint">Use Previous / Next to review one page at a time. Your text updates automatically.</p>
<div id="pptPreviewNote" class="preview-placeholder hidden">Presentation mode selected. Review the cover and key content preview below; the exported file contains the slide sequence for this profile.</div>
<div class="preview-pages" id="livePreview" aria-live="polite"><article class="preview-page" data-page="1"><div class="cover-title" id="pvTitle">PROJECT TITLE</div><div class="cover-label" id="pvType">Summer Internship Report</div><div class="cover-label" id="pvCompletedAt">Completed At<br>Organization name</div><div class="cover-label">Submitted by</div><div class="cover-name" id="pvStudent">Student name (Enrollment No.)</div><div class="cover-label" id="pvGuide">Under the guidance of: Guide name</div><div class="cover-university">JAYPEE UNIVERSITY OF ENGINEERING AND TECHNOLOGY</div><div class="cover-label" id="pvDepartment">COMPUTER SCIENCE AND ENGINEERING</div><div class="cover-label" id="pvYear">Guna, Madhya Pradesh · 2026</div><div class="preview-page-number">i</div></article><article class="preview-page" data-page="2"><div class="page-number-label">PAGE 2</div><div class="page-running-title" id="pvRunningTitle">PROJECT TITLE</div><div class="preview-chapter" id="pvChapter">1. INTRODUCTION</div><div class="preview-copy" id="pvOverview">[Your project overview will appear here.]</div><div class="preview-page-number">2</div></article><article class="preview-page" data-page="3"><div class="page-number-label">PAGE 3</div><div class="page-running-title">PROJECT DETAILS</div><div class="preview-chapter">Work / Methodology</div><div class="preview-copy" id="pvWork">[Your work details will appear here.]</div><div class="preview-page-number">3</div></article><article class="preview-page" data-page="4"><div class="page-number-label">PAGE 4</div><div class="page-running-title">PROJECT DETAILS</div><div class="preview-chapter">Results and Analysis</div><div class="preview-copy" id="pvResults">[Your results and analysis will appear here.]</div><div class="preview-page-number">4</div></article><article class="preview-page" data-page="5"><div class="page-number-label">PAGE 5</div><div class="page-running-title">PROJECT DETAILS</div><div class="preview-chapter">Individual Contribution / Work Plan</div><div class="preview-copy" id="pvContribution">[Your contribution and work plan will appear here.]</div><div class="preview-chapter">Conclusion / Future Work</div><div class="preview-copy" id="pvConclusion">[Your conclusion will appear here.]</div><div class="preview-chapter">References</div><div class="preview-copy" id="pvReferences">[Your references will appear here.]</div><div class="preview-page-number">5</div></article></div>
<div class="preview-status" id="previewStatus">Live preview · Cover and sample content</div>
</aside></div></section>
<footer>Slow-Steady · Academic document workspace · Exported files should be reviewed before submission.</footer></main>
<script>
const profileNames={internship:"Summer Internship Report",minor1:"Minor Project I Report",minor2:"Minor Project II Report",major1:"Major Project I Report",major2:"Major Project II Report",major_ppt:"Major Project PPT",minor_ppt:"Major & Minor Project PPT",internship_ppt:"Summer Internship Report PPT"};
const profileChapters={internship:"1. INTRODUCTION",minor1:"1. INTRODUCTION",minor2:"1. INTRODUCTION",major1:"1. INTRODUCTION",major2:"1. INTRODUCTION",major_ppt:"Introduction",minor_ppt:"Project Overview",internship_ppt:"Internship Overview"};
const presentationTypes=new Set(["major_ppt","minor_ppt","internship_ppt"]);
function valueOf(name,fallback){const el=document.querySelector('#builderForm [name="'+name+'"]');return el&&el.value.trim()?el.value.trim():fallback;}
function chooseDocument(type){document.getElementById("document_type").value=type;document.getElementById("dashboard").classList.remove("active");document.getElementById("workspace").classList.add("active");document.getElementById("workspaceTitle").textContent=profileNames[type];document.getElementById("workspaceSubtitle").textContent=presentationTypes.has(type)?"Enter your presentation content on the left; review the live preview on the right.":"Enter report details on the left; review the live page preview on the right.";document.getElementById("organizationField").classList.toggle("hidden",type!=="internship"&&type!=="internship_ppt");document.getElementById("pptPreviewNote").classList.toggle("hidden",!presentationTypes.has(type));document.getElementById("livePreview").classList.toggle("presentation",presentationTypes.has(type));updateLivePreview();window.scrollTo({top:0,behavior:"smooth"});}
function updateLivePreview(){
 const type=document.getElementById("document_type").value;
 const title=valueOf("project_title","PROJECT TITLE").toUpperCase();
 document.getElementById("pvTitle").textContent=title;
 document.getElementById("pvType").textContent=profileNames[type]||"Academic Report";
 document.getElementById("pvStudent").textContent=valueOf("student_name","Student name")+" ("+valueOf("enrollment","Enrollment No.")+")";
 const org=valueOf("organization","");
 document.getElementById("pvCompletedAt").innerHTML=org?"Completed At<br>"+escapeHtml(org):"Completed At<br>[Organization name]";
 const guide=valueOf("guide","");
 document.getElementById("pvGuide").textContent=guide?"Under the guidance of: "+guide:"Under the guidance of: [Guide name]";
 document.getElementById("pvDepartment").textContent=valueOf("department","Department").toUpperCase();
 document.getElementById("pvYear").textContent="Guna, Madhya Pradesh · "+valueOf("year","2026");
 document.getElementById("pvRunningTitle").textContent=title;
 document.getElementById("pvChapter").textContent=profileChapters[type]||"1. INTRODUCTION";
 document.getElementById("pvOverview").textContent=valueOf("overview","[Your project overview will appear here.]");
 document.getElementById("pvWork").textContent=valueOf("work","[Your work details will appear here.]");
 document.getElementById("pvResults").textContent=valueOf("results","[Your results and analysis will appear here.]");
 document.getElementById("pvContribution").textContent=valueOf("contribution","[Your contribution and work plan will appear here.]");
 document.getElementById("pvConclusion").textContent=valueOf("conclusion","[Your conclusion will appear here.]");
 document.getElementById("pvReferences").textContent=valueOf("references","[Your references will appear here.]");
 document.getElementById("previewStatus").textContent=(presentationTypes.has(type)?"Presentation preview · slide ratio":"A4 document preview")+" · updates while typing";showPreviewPage(0);
}
function escapeHtml(s){return s.replaceAll("&","&amp;").replaceAll("<","&lt;").replaceAll(">","&gt;").replaceAll('"',"&quot;").replaceAll("'","&#39;");}
let currentPreviewPage=0;function showPreviewPage(index){const pages=[...document.querySelectorAll("#livePreview .preview-page")];if(!pages.length)return;currentPreviewPage=Math.max(0,Math.min(index,pages.length-1));pages.forEach((p,i)=>p.classList.toggle("active",i===currentPreviewPage));document.getElementById("previewPageIndicator").textContent="Page "+(currentPreviewPage+1)+" of "+pages.length;document.getElementById("prevPreviewPage").disabled=currentPreviewPage===0;document.getElementById("nextPreviewPage").disabled=currentPreviewPage===pages.length-1;document.getElementById("prevPreviewPage").style.opacity=currentPreviewPage===0?".45":"1";document.getElementById("nextPreviewPage").style.opacity=currentPreviewPage===pages.length-1?".45":"1";}document.getElementById("prevPreviewPage").addEventListener("click",()=>showPreviewPage(currentPreviewPage-1));document.getElementById("nextPreviewPage").addEventListener("click",()=>showPreviewPage(currentPreviewPage+1));document.querySelectorAll("#builderForm input,#builderForm textarea").forEach(el=>{el.addEventListener("input",updateLivePreview);el.addEventListener("change",updateLivePreview);});
document.getElementById("builderForm").addEventListener("reset",()=>setTimeout(updateLivePreview,0));
document.querySelectorAll(".choice[data-type]").forEach(btn=>btn.addEventListener("click",()=>chooseDocument(btn.dataset.type)));
document.getElementById("backToDashboard").addEventListener("click",()=>{document.getElementById("workspace").classList.remove("active");document.getElementById("dashboard").classList.add("active");window.scrollTo({top:0,behavior:"smooth"});});
updateLivePreview();
</script></body></html>"""




def set_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = paragraph.add_run()
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); r._r.addnext(fld)

def configure_doc(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin = Inches(1); sec.bottom_margin = Inches(1)
    sec.right_margin = Inches(1); sec.left_margin = Inches(1.5)
    normal = doc.styles["Normal"]; normal.font.name = "Times New Roman"; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in [("Heading 1",16),("Heading 2",14),("Heading 3",12)]:
        st=doc.styles[name]; st.font.name="Times New Roman"; st.font.size=Pt(size); st.font.bold=True
        st.paragraph_format.keep_with_next=True
    set_page_field(sec.footer.paragraphs[0])

def centered(doc, text, size=12, bold=False):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(text); r.bold=bold; r.font.name="Times New Roman"; r.font.size=Pt(size)
    return p

def add_text(doc, text, placeholder):
    text=(text or "").strip()
    if not text:
        doc.add_paragraph("[To complete: " + placeholder + "]")
    else:
        for part in text.splitlines():
            if part.strip(): doc.add_paragraph(part.strip())

def build_report(data):
    doc=Document(); configure_doc(doc)
    typ=data["document_type"]; profile=REPORT_PROFILES[typ]
    centered(doc,data["project_title"].upper(),16,True); doc.add_paragraph()
    centered(doc,profile["label"],14,True); centered(doc,"Submitted by",14)
    centered(doc,f'{data["student_name"]} ({data["enrollment"]})'.strip(),14)
    if data["guide"]: centered(doc,"Under the guidance of: "+data["guide"],14)
    doc.add_paragraph(); centered(doc,"JAYPEE UNIVERSITY OF ENGINEERING AND TECHNOLOGY",14,True)
    centered(doc,data["department"].upper(),12); centered(doc,f'Guna, Madhya Pradesh · {data["year"]}',12)
    doc.add_page_break()
    doc.add_heading("DOCUMENT REVIEW CHECKLIST",1)
    doc.add_paragraph("This file is a draft generated from the selected profile. Replace every bracketed instruction, verify all facts and references, and check current departmental rules.")
    doc.add_page_break()
    if typ=="internship":
        doc.add_heading("DECLARATION",1)
        doc.add_paragraph(f'I hereby declare that the work reported in “{data["project_title"]}” is presented to the best of my knowledge and belief. I am responsible for ensuring the originality and accuracy of this report.')
        doc.add_paragraph("Signature of the Student\n\nPlace: ____________________\nDate: ____________________"); doc.add_page_break()
        doc.add_heading("INTERNSHIP COMPLETION CERTIFICATE",1); doc.add_paragraph("[Insert the authentic organization-issued completion certificate. Do not fabricate a certificate.]"); doc.add_page_break()
        doc.add_heading("ACKNOWLEDGEMENT",1); doc.add_paragraph("[Write a truthful acknowledgement.]"); doc.add_page_break()
        doc.add_heading("EXECUTIVE SUMMARY",1); add_text(doc,data["overview"],"one-page purpose, work done and learning outcomes"); doc.add_page_break()
        doc.add_heading("TABLE OF CONTENTS",1); doc.add_paragraph("[In Word or LibreOffice, insert/update an automatic table of contents after completing the document.]"); doc.add_page_break()
    for i,(heading,guide) in enumerate(profile["chapters"]):
        doc.add_heading(heading,1)
        if "INTRODUCTION" in heading or "PROBLEM" in heading:
            add_text(doc,data["overview"],guide)
        elif any(k in heading for k in ["WORK","IMPLEMENTATION","METHODOLOGY","SYSTEM ANALYSIS","DESCRIPTION OF WORK"]):
            add_text(doc,data["work"],guide)
        elif any(k in heading for k in ["RESULT","TESTING","ANALYSIS","EVALUATION"]):
            add_text(doc,data["results"],guide)
        elif "CONCLUSION" in heading:
            add_text(doc,data["conclusion"],guide)
        elif "REFERENCES" in heading:
            add_text(doc,data["references"],guide)
        else:
            doc.add_paragraph("[To complete: "+guide+"]")
    doc.add_heading("FINAL REVIEW CHECKLIST",1)
    checks=["Replace every placeholder with accurate content.","Insert authentic certificates and required evidence.","Verify all citations and reference details.","Check figure/table captions and references in text.","Update the table of contents and verify page numbering.","For internship reports, supplied JUET guidelines state a minimum of 40 pages after the index."]
    for c in checks: doc.add_paragraph(c,style="List Bullet")
    out=BytesIO(); doc.save(out); out.seek(0); return out

def add_slide(prs,title,body,number,total):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor(12,18,31)
    # top accent bar
    shape=slide.shapes.add_shape(1,0,0,prs.slide_width,PptInches(.12))
    shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor(110,157,255); shape.line.fill.background()
    tx=slide.shapes.add_textbox(PptInches(.7),PptInches(.55),PptInches(12),PptInches(1.1)).text_frame
    tx.word_wrap=True; p=tx.paragraphs[0]; p.text=title; p.font.name="Aptos Display"; p.font.size=PptPt(28); p.font.bold=True; p.font.color.rgb=RGBColor(239,244,255)
    bodybox=slide.shapes.add_textbox(PptInches(.85),PptInches(2),PptInches(11.5),PptInches(4.4)).text_frame
    bodybox.word_wrap=True
    lines=[x.strip() for x in (body or "").splitlines() if x.strip()]
    if not lines: lines=["[Add verified project-specific content.]"]
    for i,line in enumerate(lines[:6]):
        p=bodybox.paragraphs[0] if i==0 else bodybox.add_paragraph()
        p.text=("• " if len(lines)>1 else "")+line
        p.font.name="Aptos"; p.font.size=PptPt(21 if len(lines)<=3 else 18); p.font.color.rgb=RGBColor(211,222,241); p.space_after=PptPt(18)
    footer=slide.shapes.add_textbox(PptInches(.7),PptInches(7.05),PptInches(12),PptInches(.25)).text_frame.paragraphs[0]
    footer.text=f"Slow-Steady  ·  {number:02d} / {total:02d}"; footer.font.size=PptPt(10); footer.font.color.rgb=RGBColor(139,157,185)

def build_ppt(data):
    prs=Presentation()
    prs.slide_width=PptInches(13.333); prs.slide_height=PptInches(7.5)
    kind=data.get("presentation_type") or (data.get("document_type") if data.get("document_type") in {"major_ppt","minor_ppt","internship_ppt"} else "major_ppt")
    title=data["project_title"]
    refs=data.get("references","").strip()
    overview=data.get("overview","").strip()
    work=data.get("work","").strip()
    results=data.get("results","").strip()
    contribution=data.get("contribution","").strip()
    conclusion=data.get("conclusion","").strip()
    organization=data.get("organization","").strip()
    guide=data.get("guide","").strip()
    student=data.get("student_name","").strip()
    enrollment=data.get("enrollment","").strip()
    year=data.get("year",str(date.today().year))
    if kind=="major_ppt":
        # PPT instructions call for 1 intro + 1 problem + 1–2 literature + work plan
        # + individual contribution + 3–4 work-so-far + work remaining + conclusion + refs.
        # Exactly 12 content slides plus a cover slide = 13 slides total.
        sections=[
          ("Introduction",overview or "Add the project context, motivation and objective.",1),
          ("Problem Statement",overview or "Define the problem, intended users, scope and constraints.",1),
          ("Literature Review","Summarize verified prior work and the gap addressed by this project.",1),
          ("Work Plan",contribution or "Add the weekly-report-aligned milestones and dates.",1),
          ("Individual Contribution",contribution or f"State the verified contribution of {student or 'each team member'}.",1),
          ("Work Done So Far",work or "Add evidence-backed work completed so far.",4),
          ("Work to be Done","Add remaining tasks, milestones, owners and risks.",1),
          ("Conclusion",conclusion or "Summarize only conclusions supported by current evidence.",1),
          ("References",refs or "Add verified references in IEEE format.",1)
        ]
        slide_items=[]
        for heading,body,count in sections:
            for n in range(count):
                text=body
                if heading=="Work Done So Far" and n==1: text=results or "Add verified implementation or methodology evidence."
                if heading=="Work Done So Far" and n==2: text="Add testing, experiments or analysis with actual evidence."
                if heading=="Work Done So Far" and n==3: text="Add a distinct design, evaluation or results summary."
                slide_items.append((heading if count==1 else f"{heading} ({n+1})",text))
    elif kind=="minor_ppt":
        slide_items=[
          ("Project Title",f"{title}\n{student} · {enrollment}\nMinor Project Presentation · {year}"),
          ("Introduction",overview or "Introduce the project domain, motivation and intended users."),
          ("Problem Statement",overview or "Define the specific problem, scope and objectives."),
          ("Literature Review", "Summarize verified related work and identify the gap."),
          ("Objectives", "List measurable objectives and expected deliverables."),
          ("Proposed Methodology",work or "Describe the proposed approach, tools, architecture and workflow."),
          ("System Design / Research Method", "Add architecture/UML for application-based work, or dataset and methodology for research-based work."),
          ("Work Completed",work or "Describe the actual progress completed so far."),
          ("Work Plan",contribution or "Add milestones and timeline aligned with the project plan."),
          ("Expected Outcomes / Initial Results",results or "Describe expected outcomes or verified initial results; do not invent results."),
          ("Conclusion and Next Steps",conclusion or "Summarize the current status and the next milestone."),
          ("References",refs or "Add the sources cited in the presentation.")
        ]
    else:
        slide_items=[
          ("Internship Overview",f"{title}\n{student} · {enrollment}\n{organization or 'Organization name'} · {year}"),
          ("Organization Profile",organization or "Add verified details about the organization, team and domain."),
          ("Internship Objectives",overview or "Describe the internship objectives and expected learning outcomes."),
          ("Problem / Project Scope",overview or "Explain the assigned problem, users and scope."),
          ("Tools and Technologies",work or "List the technologies and tools actually used."),
          ("Work Completed",work or "Describe the work completed during the internship."),
          ("Implementation / Methodology",work or "Explain the approach, design and implementation steps."),
          ("Results and Learning",results or "Present verified outcomes, tests and lessons learned."),
          ("Challenges and Solutions","Describe real challenges encountered and how they were addressed."),
          ("Conclusion and Future Work",conclusion or "Summarize outcomes and possible next steps."),
          ("References",refs or "Add sources cited in the presentation.")
        ]
    total=len(slide_items)+1
    add_slide(prs,title, f"{student}\n{enrollment}\n{organization if kind=='internship_ppt' else 'Jaypee University of Engineering and Technology'}\n{year}",1,total)
    for i,(heading,body) in enumerate(slide_items,2):
        add_slide(prs,heading,body,i,total)
    out=BytesIO(); prs.save(out); out.seek(0); return out

@app.get("/api/self-test")
def self_test():
    """Generate and reopen sample outputs to verify the deployed document libraries."""
    sample = {
        "document_type": "internship", "project_title": "Automated Output Test",
        "student_name": "Test Student", "enrollment": "TEST-001",
        "organization": "Test Organization", "guide": "Test Guide",
        "department": "Computer Science and Engineering", "year": "2026",
        "overview": "Sample overview for automated output verification.",
        "work": "Sample work and methodology.", "results": "Sample test results.",
        "contribution": "Sample contribution and plan.", "conclusion": "Sample conclusion.",
        "references": "[1] Sample reference for output verification."
    }
    checks = {}
    errors = []
    try:
        doc_bytes = build_report(sample).getvalue()
        parsed_doc = Document(BytesIO(doc_bytes))
        checks["docx_generated_and_reopened"] = bool(parsed_doc.paragraphs)
        checks["docx_a4_and_margins"] = (
            abs(parsed_doc.sections[0].page_width.inches - 8.27) < 0.02
            and abs(parsed_doc.sections[0].page_height.inches - 11.69) < 0.02
            and abs(parsed_doc.sections[0].left_margin.inches - 1.5) < 0.01
            and abs(parsed_doc.sections[0].right_margin.inches - 1.0) < 0.01
        )
        checks["docx_times_new_roman_12pt"] = (
            parsed_doc.styles["Normal"].font.name == "Times New Roman"
            and parsed_doc.styles["Normal"].font.size.pt == 12
        )
        checks["all_five_report_profiles_generate"] = True
        for profile in REPORT_PROFILES:
            candidate = dict(sample, document_type=profile)
            reopened = Document(BytesIO(build_report(candidate).getvalue()))
            if not reopened.paragraphs:
                checks["all_five_report_profiles_generate"] = False
        ppt_bytes = build_ppt(sample).getvalue()
        parsed_ppt = Presentation(BytesIO(ppt_bytes))
        checks["pptx_generated_and_reopened"] = len(parsed_ppt.slides) == 13
        ppt_text = "\\n".join(
            shape.text for slide in parsed_ppt.slides
            for shape in slide.shapes if shape.has_text_frame
        )
        checks["pptx_required_sections"] = all(
            heading in ppt_text for heading in (
                "Problem Statement", "Literature Review", "Work Plan",
                "Individual Contribution", "Work Done So Far",
                "Work to be Done", "Conclusion", "References"
            )
        )
    except Exception as exc:
        errors.append(f"{type(exc).__name__}: {exc}")
    passed = bool(checks) and all(checks.values()) and not errors
    return JSONResponse(
        {"status": "passed" if passed else "failed", "checks": checks, "errors": errors},
        status_code=200 if passed else 500
    )

@app.get("/",response_class=HTMLResponse)
def home(): return HTMLResponse(PAGE)

@app.get("/health")
def health(): return {"status":"ok","service":"slow-steady","version":"0.2.0","profiles":list(REPORT_PROFILES),"pptx":True}

@app.get("/api/profiles")
def profiles():
    return {"reports":{k:{"label":v["label"],"chapters":[x[0] for x in v["chapters"]]} for k,v in REPORT_PROFILES.items()},
            "major_ppt":[{"section":t,"slides":n,"guidance":g} for t,n,g in PPT_SECTIONS],
            "note":"Minor/Major report headings are initial starter profiles and must be checked against the relevant official guideline."}

@app.post("/generate")
def generate(document_type:str=Form(...),project_title:str=Form(...),student_name:str=Form(...),enrollment:str=Form(""),organization:str=Form(""),guide:str=Form(""),department:str=Form("Computer Science and Engineering"),year:str=Form(str(date.today().year)),overview:str=Form(""),work:str=Form(""),results:str=Form(""),contribution:str=Form(""),conclusion:str=Form(""),references:str=Form("")):
    data={k:v.strip() if isinstance(v,str) else v for k,v in locals().copy().items() if k not in {"document_type"}}
    data["document_type"]=document_type
    if not project_title.strip() or not student_name.strip(): return JSONResponse({"error":"Project title and student/presenter are required."},status_code=422)
    if document_type in {"major_ppt","minor_ppt","internship_ppt"}:
        data["presentation_type"]=document_type
        out=build_ppt(data); filename=f"Slow-Steady-{document_type.replace('_','-')}.pptx"; media="application/vnd.openxmlformats-officedocument.presentationml.presentation"
    elif document_type in REPORT_PROFILES:
        out=build_report(data); filename=f'Slow-Steady-{document_type}-Report.docx'; media="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else: return JSONResponse({"error":"Unknown document profile."},status_code=422)
    return StreamingResponse(out,media_type=media,headers={"Content-Disposition":f'attachment; filename="{filename}"'})

@app.post("/api/validate")
def validate(document_type:str=Form(...),project_title:str=Form(""),student_name:str=Form(""),overview:str=Form(""),work:str=Form(""),results:str=Form(""),references:str=Form("")):
    issues=[]
    for key,label in [("project_title","Project title"),("student_name","Student/presenter name")]:
        if not {"project_title": project_title, "student_name": student_name}[key].strip(): issues.append({"severity":"error","rule":f"required_{key}","message":f"{label} is missing."})
    if not overview.strip(): issues.append({"severity":"warning","rule":"overview_present","message":"Project overview/problem statement is empty."})
    if not work.strip(): issues.append({"severity":"warning","rule":"work_present","message":"Work completed/methodology is empty."})
    if not results.strip(): issues.append({"severity":"warning","rule":"results_present","message":"Results/analysis is empty."})
    if not references.strip(): issues.append({"severity":"warning","rule":"references_present","message":"No references supplied; add verified citations where used."})
    if document_type=="major_ppt":
        return {"profile":document_type,"checks":{"required_sections":len(PPT_SECTIONS),"expected_slide_count":sum(x[1] for x in PPT_SECTIONS),"ai_generated_images":False},"issues":issues,"passed":not any(i["severity"]=="error" for i in issues)}
    return {"profile":document_type,"checks":{"required_heading_profile":document_type in REPORT_PROFILES,"a4_page_setup":True,"times_new_roman_12pt_body":True,"line_spacing_1_15":True},"issues":issues,"passed":not any(i["severity"]=="error" for i in issues)}
